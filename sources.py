import ipaddress
import socket
from io import BytesIO
from urllib.parse import urljoin, urlparse

import httpx
from bs4 import BeautifulSoup
from docx import Document
from pypdf import PdfReader
from pptx import Presentation

from config import MAX_SOURCE_CHARS
from logging_config import configure_logging


logger = configure_logging().getChild("sources")


def _validate_public_url(url: str) -> None:
    parsed = urlparse(url)
    if parsed.scheme not in {"http", "https"} or not parsed.hostname:
        raise ValueError("Only public http:// and https:// URLs are supported.")
    if parsed.username or parsed.password:
        raise ValueError("URLs containing credentials are not supported.")
    try:
        addresses = socket.getaddrinfo(parsed.hostname, parsed.port or (443 if parsed.scheme == "https" else 80))
    except socket.gaierror as error:
        raise ValueError("The source URL could not be resolved.") from error
    for address in addresses:
        ip = ipaddress.ip_address(address[4][0].split("%", 1)[0])
        if not ip.is_global:
            raise ValueError("Private or local network URLs are not supported.")


def fetch_public_url(url: str) -> str:
    current_url = url
    headers = {"User-Agent": "SlideGenerator/1.0 (source reader)"}
    with httpx.Client(timeout=15, follow_redirects=False, headers=headers) as client:
        for _ in range(4):
            _validate_public_url(current_url)
            response = client.get(current_url)
            if response.status_code in {301, 302, 303, 307, 308}:
                location = response.headers.get("location")
                if not location:
                    break
                current_url = urljoin(current_url, location)
                continue
            response.raise_for_status()
            if len(response.content) > 3 * 1024 * 1024:
                raise ValueError("The webpage is too large to process (3 MB maximum).")
            content_type = response.headers.get("content-type", "").lower()
            if not any(kind in content_type for kind in ("text/html", "text/plain", "application/xhtml")):
                raise ValueError("The URL must point to a readable webpage or plain-text document.")
            if "text/html" in content_type or "xhtml" in content_type:
                soup = BeautifulSoup(response.text, "html.parser")
                for element in soup(["script", "style", "nav", "footer", "header", "noscript"]):
                    element.decompose()
                return soup.get_text(" ", strip=True)[:MAX_SOURCE_CHARS]
            return response.text[:MAX_SOURCE_CHARS]
    raise ValueError("The source URL redirected too many times.")


def collect_source_text(
    input_text: str,
    source_url: str,
    request_id: str = "",
    source_id_offset: int = 0,
) -> str:
    urls = [line.strip() for line in source_url.splitlines() if line.strip()]
    headers = [
        f"[S{source_id_offset + index}] Web source: {url}\n"
        for index, url in enumerate(urls, start=1)
    ]
    available_chars = max(0, MAX_SOURCE_CHARS - sum(len(header) for header in headers))
    section_count = len(urls) + bool(input_text.strip())
    section_budget = available_chars // max(1, section_count)
    parts = [input_text.strip()[:section_budget]] if input_text.strip() else []
    for index, (url, header) in enumerate(zip(urls, headers), start=1):
        logger.info("Request %s: reading web source %d/%d", request_id, index, len(urls))
        parts.append(f"{header}{fetch_public_url(url)[:section_budget]}")
    combined = "\n\n".join(parts)
    if not combined.strip():
        raise ValueError("No readable text was found in the provided inputs.")
    return combined[:MAX_SOURCE_CHARS]


def extract_document(filename: str, content: bytes) -> str:
    suffix = filename.lower().rsplit(".", 1)[-1] if "." in filename else ""
    if suffix in {"txt", "md"}:
        return content.decode("utf-8-sig", errors="replace")
    if suffix == "pdf":
        reader = PdfReader(BytesIO(content))
        if reader.is_encrypted:
            raise ValueError(f"{filename} is password-protected.")
        return "\n".join(page.extract_text() or "" for page in reader.pages)
    if suffix == "docx":
        document = Document(BytesIO(content))
        return "\n".join(paragraph.text for paragraph in document.paragraphs)
    if suffix == "pptx":
        presentation = Presentation(BytesIO(content))
        slide_text = []
        for slide_number, slide in enumerate(presentation.slides, start=1):
            text = [shape.text.strip() for shape in slide.shapes if shape.has_text_frame and shape.text.strip()]
            slide_text.append(f"Slide {slide_number}:\n" + "\n".join(text))
        return "\n\n".join(slide_text)
    raise ValueError(f"{filename}: use PPTX, PDF, DOCX, TXT, or Markdown files.")