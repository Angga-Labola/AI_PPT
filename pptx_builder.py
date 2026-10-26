from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR
from pptx.util import Inches, Pt


INK = RGBColor(25, 43, 37)
PAPER = RGBColor(247, 247, 239)
GREEN = RGBColor(143, 213, 93)
CORAL = RGBColor(232, 119, 96)
MUTED = RGBColor(124, 143, 131)


def _set_background(slide, color: RGBColor) -> None:
    fill = slide.background.fill
    fill.solid()
    fill.fore_color.rgb = color


def _add_text(slide, text: str, left, top, width, height, *, size, color, bold=False):
    shape = slide.shapes.add_textbox(left, top, width, height)
    frame = shape.text_frame
    frame.clear()
    frame.word_wrap = True
    frame.vertical_anchor = MSO_ANCHOR.MIDDLE
    paragraph = frame.paragraphs[0]
    paragraph.text = text
    paragraph.font.name = "Aptos"
    paragraph.font.size = Pt(size)
    paragraph.font.bold = bold
    paragraph.font.color.rgb = color
    return shape


def build_presentation(deck, path: Path, source_catalog: dict[str, str] | None = None) -> None:
    presentation = Presentation()
    presentation.slide_width = Inches(13.333)
    presentation.slide_height = Inches(7.5)
    blank = presentation.slide_layouts[6]

    cover = presentation.slides.add_slide(blank)
    _set_background(cover, INK)
    accent = cover.shapes.add_shape(
        MSO_SHAPE.RECTANGLE, Inches(0.75), Inches(0.75), Inches(0.14), Inches(5.9)
    )
    accent.fill.solid()
    accent.fill.fore_color.rgb = GREEN
    accent.line.fill.background()
    _add_text(cover, "FIELD NOTES  /  BRIEFING", Inches(1.25), Inches(1.1), Inches(8), Inches(0.4), size=12, color=GREEN, bold=True)
    _add_text(cover, deck.title, Inches(1.2), Inches(2.0), Inches(10.8), Inches(2.3), size=36, color=PAPER, bold=True)
    if deck.subtitle:
        _add_text(cover, deck.subtitle, Inches(1.25), Inches(4.55), Inches(9.5), Inches(0.9), size=18, color=MUTED)
    dot = cover.shapes.add_shape(
        MSO_SHAPE.OVAL, Inches(11.6), Inches(5.65), Inches(0.7), Inches(0.7)
    )
    dot.fill.solid()
    dot.fill.fore_color.rgb = CORAL
    dot.line.fill.background()

    for index, item in enumerate(deck.slides, start=1):
        slide = presentation.slides.add_slide(blank)
        _set_background(slide, PAPER)
        marker = slide.shapes.add_shape(
            MSO_SHAPE.RECTANGLE, Inches(0.72), Inches(0.72), Inches(0.72), Inches(0.11)
        )
        marker.fill.solid()
        marker.fill.fore_color.rgb = GREEN if index % 2 else CORAL
        marker.line.fill.background()
        _add_text(slide, f"{index:02d}  /  {len(deck.slides):02d}", Inches(11.3), Inches(0.55), Inches(1.2), Inches(0.4), size=11, color=MUTED, bold=True)
        _add_text(slide, item.title, Inches(0.85), Inches(1.15), Inches(11.3), Inches(1.0), size=27, color=INK, bold=True)
        body = slide.shapes.add_textbox(Inches(1.0), Inches(2.35), Inches(11.1), Inches(4.2))
        frame = body.text_frame
        frame.clear()
        frame.word_wrap = True
        for bullet_index, bullet in enumerate(item.bullets):
            paragraph = frame.paragraphs[0] if bullet_index == 0 else frame.add_paragraph()
            paragraph.text = f"{bullet_index + 1:02d}    {bullet}"
            paragraph.font.name = "Aptos"
            paragraph.font.size = Pt(20)
            paragraph.font.color.rgb = INK
            paragraph.space_after = Pt(18)
            paragraph.level = 0
        citation = slide.shapes.add_textbox(
            Inches(0.9), Inches(6.72), Inches(11.5), Inches(0.48)
        )
        citation_frame = citation.text_frame
        citation_frame.clear()
        citation_frame.word_wrap = True
        citation_paragraph = citation_frame.paragraphs[0]
        citation_label = citation_paragraph.add_run()
        citation_label.text = "Sources: "
        citation_label.font.name = "Aptos"
        citation_label.font.size = Pt(10)
        citation_label.font.color.rgb = MUTED
        for source_index, source_id in enumerate(item.sources):
            if source_index:
                separator = citation_paragraph.add_run()
                separator.text = ", "
                separator.font.size = Pt(10)
                separator.font.color.rgb = MUTED
            source_run = citation_paragraph.add_run()
            source_name = (source_catalog or {}).get(source_id, "")
            source_run.text = (
                f"{source_id} ({source_name[:48]})" if source_name and not source_name.startswith(("http://", "https://"))
                else source_id
            )
            source_run.font.name = "Aptos"
            source_run.font.size = Pt(10)
            source_run.font.color.rgb = MUTED
            if source_name.startswith(("http://", "https://")):
                source_run.hyperlink.address = source_name
        footer = slide.shapes.add_shape(
            MSO_SHAPE.RECTANGLE, Inches(0), Inches(7.34), Inches(13.333), Inches(0.16)
        )
        footer.fill.solid()
        footer.fill.fore_color.rgb = INK
        footer.line.fill.background()

    path.parent.mkdir(parents=True, exist_ok=True)
    presentation.save(path)