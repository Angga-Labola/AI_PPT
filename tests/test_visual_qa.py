import tempfile
import unittest
from pathlib import Path

from pptx import Presentation
from pptx.util import Inches

from pptx_builder import build_presentation
from visual_qa import VisualLayoutError, render_presentation_previews

from deck_generator import SlideDeck


def _deck(title: str = "Visual QA Test Deck") -> SlideDeck:
    return SlideDeck.model_validate(
        {
            "title": title,
            "subtitle": "A test presentation",
            "slides": [
                {
                    "title": "A concise slide title",
                    "bullets": [
                        {"text": f"Grounded content point {index}.", "sources": ["S1"]}
                        for index in range(1, 4)
                    ],
                }
                for _ in range(5)
            ],
        }
    )


class VisualQATests(unittest.TestCase):
    def test_renders_previews_for_a_valid_deck(self) -> None:
        with tempfile.TemporaryDirectory() as temporary:
            root = Path(temporary)
            presentation_path = root / "valid.pptx"
            preview_dir = root / "previews"
            build_presentation(_deck(), presentation_path, {"S1": "Test source"})
            slide_count = render_presentation_previews(presentation_path, preview_dir)
            self.assertEqual(slide_count, 6)
            self.assertEqual(len(list(preview_dir.glob("slide-*.png"))), 6)

    def test_rejects_a_slide_title_that_overflows(self) -> None:
        with tempfile.TemporaryDirectory() as temporary:
            root = Path(temporary)
            deck = _deck()
            presentation_path = root / "overflow.pptx"
            build_presentation(deck, presentation_path, {"S1": "Test source"})
            presentation = Presentation(presentation_path)
            title_shape = next(
                shape
                for shape in presentation.slides[1].shapes
                if shape.has_text_frame and shape.text == deck.slides[0].title
            )
            title_shape.height = Inches(0.1)
            presentation.save(presentation_path)
            with self.assertRaises(VisualLayoutError):
                render_presentation_previews(presentation_path, root / "previews")


if __name__ == "__main__":
    unittest.main()
