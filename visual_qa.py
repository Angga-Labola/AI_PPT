import logging
import os
from functools import lru_cache
from pathlib import Path

from PIL import Image, ImageDraw, ImageFont
from pptx import Presentation
from pptx.enum.dml import MSO_COLOR_TYPE
from pptx.enum.shapes import MSO_SHAPE, MSO_SHAPE_TYPE


logger = logging.getLogger("slides_generator.visual_qa")
PREVIEW_WIDTH = 1280
PREVIEW_HEIGHT = 720
BODY_WIDTH_POINTS = 10.9 * 72
BODY_HEIGHT_POINTS = 4.1 * 72
BODY_FONT_SIZES = tuple(range(20, 13, -1))
_font_fallback_logged = False


class VisualLayoutError(ValueError):
    pass


def _font_path(bold: bool) -> str | None:
    font_dir = Path(os.environ.get("WINDIR", r"C:\Windows")) / "Fonts"
    candidates = (
        ("Aptos-Bold.ttf", "arialbd.ttf", "Aptos.ttf", "arial.ttf")
        if bold
        else ("Aptos.ttf", "arial.ttf")
    )
    for name in candidates:
        candidate = font_dir / name
        if candidate.is_file():
            return str(candidate)
    linux_font_dir = Path("/usr/share/fonts/truetype/dejavu")
    linux_font = linux_font_dir / (
        "DejaVuSans-Bold.ttf" if bold else "DejaVuSans.ttf"
    )
    return str(linux_font) if linux_font.is_file() else None


@lru_cache(maxsize=32)
def _font(size: int, bold: bool = False) -> ImageFont.FreeTypeFont | ImageFont.ImageFont:
    global _font_fallback_logged
    path = _font_path(bold)
    if path:
        return ImageFont.truetype(path, size=size)
    if not _font_fallback_logged:
        logger.warning(
            "No preview font found; using Pillow's default font for slide layout checks."
        )
        _font_fallback_logged = True
    return ImageFont.load_default(size=size)


def _rgb(color, default: tuple[int, int, int]) -> tuple[int, int, int]:
    if color.type == MSO_COLOR_TYPE.RGB and color.rgb:
        return tuple(color.rgb)
    return default


def _wrap_text(
    draw: ImageDraw.ImageDraw,
    text: str,
    font: ImageFont.FreeTypeFont | ImageFont.ImageFont,
    max_width: int,
) -> list[str]:
    lines = []
    for paragraph in text.splitlines() or [""]:
        line = ""
        for word in paragraph.split():
            candidate = f"{line} {word}".strip()
            if draw.textbbox((0, 0), candidate, font=font)[2] <= max_width:
                line = candidate
            elif line:
                lines.append(line)
                line = word
            else:
                fragment = ""
                for character in word:
                    candidate = fragment + character
                    if fragment and draw.textbbox((0, 0), candidate, font=font)[2] > max_width:
                        lines.append(fragment)
                        fragment = character
                    else:
                        fragment = candidate
                line = fragment
        lines.append(line)
    return lines


def body_font_size(bullets: list[str]) -> int:
    for size in BODY_FONT_SIZES:
        font = _font(round(size * 96 / 72))
        canvas = Image.new("RGB", (PREVIEW_WIDTH, PREVIEW_HEIGHT))
        draw = ImageDraw.Draw(canvas)
        line_height = draw.textbbox((0, 0), "Ag", font=font)[3] + 4
        total_height = 0
        for bullet in bullets:
            total_height += len(
                _wrap_text(draw, bullet, font, round(BODY_WIDTH_POINTS * 96 / 72))
            ) * line_height
        total_height += len(bullets) * round(18 * 96 / 72)
        if total_height <= round(BODY_HEIGHT_POINTS * 96 / 72):
            return size
    raise VisualLayoutError(
        "Slide content is too dense to fit the content area legibly. "
        "Shorten the bullets or split the content across slides."
    )


def _fits_text(text: str, size: int, width_points: float, height_points: float) -> bool:
    font = _font(round(size * 96 / 72))
    canvas = Image.new("RGB", (PREVIEW_WIDTH, PREVIEW_HEIGHT))
    draw = ImageDraw.Draw(canvas)
    width = round(width_points * 96 / 72)
    height = round(height_points * 96 / 72)
    lines = _wrap_text(draw, text, font, width)
    line_height = draw.textbbox((0, 0), "Ag", font=font)[3] + 4
    return len(lines) * line_height <= height


def validate_deck_text_fit(deck) -> None:
    if not _fits_text(deck.title, 36, 10.6 * 72, 2.2 * 72):
        raise VisualLayoutError("The presentation title does not fit the cover layout.")
    if deck.subtitle and not _fits_text(deck.subtitle, 18, 9.3 * 72, 0.8 * 72):
        raise VisualLayoutError("The presentation subtitle does not fit the cover layout.")
    for slide_number, slide in enumerate(deck.slides, start=1):
        if not _fits_text(slide.title, 27, 11.1 * 72, 0.9 * 72):
            raise VisualLayoutError(
                f"Slide {slide_number} has a title that does not fit the title layout."
            )


def _render_text_shape(
    draw: ImageDraw.ImageDraw,
    shape,
    scale_x: float,
    scale_y: float,
    slide_number: int,
) -> None:
    frame = shape.text_frame
    left = round((shape.left + frame.margin_left) * scale_x)
    top = round((shape.top + frame.margin_top) * scale_y)
    right = round((shape.left + shape.width - frame.margin_right) * scale_x)
    bottom = round((shape.top + shape.height - frame.margin_bottom) * scale_y)
    y = top
    for paragraph in frame.paragraphs:
        paragraph_text = paragraph.text
        if not paragraph_text:
            continue
        font_size = paragraph.font.size
        bold = paragraph.font.bold or False
        for run in paragraph.runs:
            if run.font.size and (font_size is None or run.font.size > font_size):
                font_size = run.font.size
            bold = bold or bool(run.font.bold)
        pixels = round((font_size.pt if font_size else 18) * scale_y)
        font = _font(max(8, pixels), bold)
        color = _rgb(paragraph.font.color, (25, 43, 37))
        lines = _wrap_text(draw, paragraph_text, font, max(1, right - left))
        line_height = draw.textbbox((0, 0), "Ag", font=font)[3] + 4
        for line in lines:
            if y + line_height > bottom:
                raise VisualLayoutError(
                    f"Slide {slide_number} has text that overflows its text box."
                )
            draw.text((left, y), line, font=font, fill=color)
            y += line_height
        if paragraph.space_after:
            y += round(paragraph.space_after.pt * scale_y)
    if y > bottom:
        raise VisualLayoutError(
            f"Slide {slide_number} has text that overflows its text box."
        )


def render_presentation_previews(presentation_path: Path, preview_dir: Path) -> int:
    presentation = Presentation(presentation_path)
    scale_x = PREVIEW_WIDTH / presentation.slide_width
    scale_y = PREVIEW_HEIGHT / presentation.slide_height
    preview_dir.mkdir(parents=True, exist_ok=True)

    for slide_number, slide in enumerate(presentation.slides, start=1):
        try:
            background = _rgb(slide.background.fill.fore_color, (247, 247, 239))
        except (AttributeError, ValueError):
            background = (247, 247, 239)
        image = Image.new("RGB", (PREVIEW_WIDTH, PREVIEW_HEIGHT), background)
        draw = ImageDraw.Draw(image)
        for shape in slide.shapes:
            if (
                shape.left < 0
                or shape.top < 0
                or shape.left + shape.width > presentation.slide_width
                or shape.top + shape.height > presentation.slide_height
            ):
                raise VisualLayoutError(
                    f"Slide {slide_number} contains a shape outside the slide bounds."
                )
            box = (
                round(shape.left * scale_x),
                round(shape.top * scale_y),
                round((shape.left + shape.width) * scale_x),
                round((shape.top + shape.height) * scale_y),
            )
            if shape.shape_type == MSO_SHAPE_TYPE.AUTO_SHAPE:
                try:
                    fill = _rgb(shape.fill.fore_color, (247, 247, 239))
                except (AttributeError, ValueError):
                    fill = (247, 247, 239)
                if shape.auto_shape_type == MSO_SHAPE.OVAL:
                    draw.ellipse(box, fill=fill)
                else:
                    draw.rectangle(box, fill=fill)
            if shape.has_text_frame and shape.text.strip():
                _render_text_shape(draw, shape, scale_x, scale_y, slide_number)
        image.save(preview_dir / f"slide-{slide_number:02d}.png")
    return len(presentation.slides)


def validate_slide_content(bullets: list[str]) -> int:
    return body_font_size(bullets)
